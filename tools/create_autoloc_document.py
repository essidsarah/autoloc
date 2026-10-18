from pathlib import Path
from PIL import Image, ImageDraw, ImageFont
from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.section import WD_SECTION
from docx.shared import Inches, Pt, RGBColor
from docx.oxml import OxmlElement
from docx.oxml.ns import qn


ROOT = Path(__file__).resolve().parents[1]
OUTPUT = ROOT / "livrables" / "AutoLoc_Couches_Repository_Service.docx"
CAPTURES = ROOT / "livrables" / "captures_code"
PACKAGE = "src/main/java/tn/esprit/autoloc"
ENTITIES = [
    "Agence", "Client", "Contrat", "Employe", "Equipement",
    "Maintenance", "Paiement", "Reservation", "Vehicule",
]


def font(size, bold=False):
    candidates = [
        "C:/Windows/Fonts/consolab.ttf" if bold else "C:/Windows/Fonts/consola.ttf",
        "C:/Windows/Fonts/lucon.ttf",
    ]
    for candidate in candidates:
        if Path(candidate).exists():
            return ImageFont.truetype(candidate, size)
    return ImageFont.load_default()


def capture(source: Path, target: Path):
    lines = source.read_text(encoding="utf-8").splitlines()
    code_font = font(25)
    title_font = font(24, bold=True)
    margin = 34
    line_height = 35
    width = 1800
    height = 105 + margin * 2 + max(1, len(lines)) * line_height
    image = Image.new("RGB", (width, height), "#161b22")
    draw = ImageDraw.Draw(image)
    draw.rounded_rectangle((0, 0, width, 105), radius=0, fill="#24292f")
    draw.ellipse((30, 35, 48, 53), fill="#ff5f57")
    draw.ellipse((58, 35, 76, 53), fill="#febc2e")
    draw.ellipse((86, 35, 104, 53), fill="#28c840")
    draw.text((130, 29), source.name, font=title_font, fill="#e6edf3")
    number_width = 70
    for index, line in enumerate(lines, start=1):
        y = 105 + margin + (index - 1) * line_height
        draw.text((margin, y), f"{index:>2}", font=code_font, fill="#7d8590")
        draw.text((margin + number_width, y), line.expandtabs(4), font=code_font, fill="#e6edf3")
    target.parent.mkdir(parents=True, exist_ok=True)
    image.save(target)


def set_cell_shading(cell, fill):
    tc_pr = cell._tc.get_or_add_tcPr()
    shading = OxmlElement("w:shd")
    shading.set(qn("w:fill"), fill)
    tc_pr.append(shading)


def set_repeat_table_header(row):
    tr_pr = row._tr.get_or_add_trPr()
    element = OxmlElement("w:tblHeader")
    element.set(qn("w:val"), "true")
    tr_pr.append(element)


def set_run_font(run, size=None, bold=None, color=None):
    run.font.name = "Aptos"
    run._element.rPr.rFonts.set(qn("w:ascii"), "Aptos")
    run._element.rPr.rFonts.set(qn("w:hAnsi"), "Aptos")
    if size:
        run.font.size = Pt(size)
    if bold is not None:
        run.bold = bold
    if color:
        run.font.color.rgb = RGBColor(*color)


def add_caption(doc, text):
    paragraph = doc.add_paragraph()
    paragraph.paragraph_format.space_before = Pt(4)
    paragraph.paragraph_format.space_after = Pt(3)
    run = paragraph.add_run(text)
    set_run_font(run, size=9, bold=True, color=(0, 0, 0))


def add_capture(doc, label, source: Path):
    target = CAPTURES / f"{source.stem}.png"
    capture(source, target)
    add_caption(doc, label)
    paragraph = doc.add_paragraph()
    paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
    paragraph.paragraph_format.space_after = Pt(5)
    paragraph.add_run().add_picture(str(target), width=Inches(6.7))


def build_document():
    OUTPUT.parent.mkdir(parents=True, exist_ok=True)
    document = Document()
    section = document.sections[0]
    section.top_margin = Inches(0.65)
    section.bottom_margin = Inches(0.65)
    section.left_margin = Inches(0.7)
    section.right_margin = Inches(0.7)

    styles = document.styles
    styles["Normal"].font.name = "Aptos"
    styles["Normal"]._element.rPr.rFonts.set(qn("w:ascii"), "Aptos")
    styles["Normal"]._element.rPr.rFonts.set(qn("w:hAnsi"), "Aptos")
    styles["Normal"].font.size = Pt(10)
    for style_name in ["Title", "Heading 1", "Heading 2"]:
        style = styles[style_name]
        style.font.name = "Aptos"
        style._element.rPr.rFonts.set(qn("w:ascii"), "Aptos")
        style._element.rPr.rFonts.set(qn("w:hAnsi"), "Aptos")
        style.font.color.rgb = RGBColor(0, 0, 0)

    title = document.add_paragraph(style="Title")
    title.alignment = WD_ALIGN_PARAGRAPH.CENTER
    title.add_run("AutoLoc Couches Repository et Service")
    intro = document.add_paragraph()
    intro.alignment = WD_ALIGN_PARAGRAPH.CENTER
    intro.add_run("Dossier de dépôt des interfaces Repository Service et des implémentations CRUD")
    intro.runs[0].italic = True
    document.add_paragraph("")

    document.add_paragraph("Contenu", style="Heading 1")
    paragraph = document.add_paragraph()
    paragraph.add_run("Ce document présente les couches créées pour les neuf entités du projet AutoLoc. ")
    paragraph.add_run("Chaque service utilise l’injection par constructeur de son repository et porte l’annotation ").bold = False
    annotation = paragraph.add_run("@Service")
    annotation.bold = True
    paragraph.add_run(". Chaque repository étend JpaRepository et porte l’annotation ")
    repository_annotation = paragraph.add_run("@Repository")
    repository_annotation.bold = True
    paragraph.add_run(".")

    document.add_paragraph("Synthèse", style="Heading 1")
    table = document.add_table(rows=1, cols=4)
    table.style = "Table Grid"
    headers = ["Entité", "Repository", "Service", "Implémentation"]
    for cell, header in zip(table.rows[0].cells, headers):
        cell.text = header
        set_cell_shading(cell, "1F4E78")
        for run in cell.paragraphs[0].runs:
            set_run_font(run, size=9, bold=True, color=(255, 255, 255))
    set_repeat_table_header(table.rows[0])
    for entity in ENTITIES:
        cells = table.add_row().cells
        values = [entity, f"{entity}Repository", f"{entity}Service", f"{entity}ServiceImpl"]
        for cell, value in zip(cells, values):
            cell.text = value
            for run in cell.paragraphs[0].runs:
                set_run_font(run, size=9)

    document.add_page_break()
    for position, entity in enumerate(ENTITIES):
        document.add_paragraph(entity, style="Heading 1")
        base = ROOT / PACKAGE
        add_capture(document, f"Repository  {entity}Repository", base / "repositories" / f"{entity}Repository.java")
        add_capture(document, f"Service  {entity}Service", base / "services" / f"{entity}Service.java")
        document.add_page_break()
        document.add_paragraph(f"Implémentation {entity}", style="Heading 1")
        add_capture(document, f"Implémentation  {entity}ServiceImpl", base / "services" / "impl" / f"{entity}ServiceImpl.java")
        if position < len(ENTITIES) - 1:
            document.add_page_break()

    for section in document.sections:
        footer = section.footer.paragraphs[0]
        footer.alignment = WD_ALIGN_PARAGRAPH.CENTER
        run = footer.add_run("AutoLoc  Couches Repository et Service")
        set_run_font(run, size=8, color=(80, 80, 80))
    document.save(OUTPUT)
    print(OUTPUT)


if __name__ == "__main__":
    build_document()

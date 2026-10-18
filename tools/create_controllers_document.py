from pathlib import Path
from PIL import Image, ImageDraw, ImageFont
from docx import Document
from docx.enum.section import WD_ORIENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.shared import Inches

root = Path(__file__).resolve().parents[1]
out = root / 'livrables' / 'AutoLoc_Controllers_et_Tests_Postman.docx'
shots = root / 'livrables' / 'captures_controllers'
names = ['Agence','Client','Contrat','Employe','Equipement','Maintenance','Paiement','Reservation','Vehicule']
font = ImageFont.truetype('C:/Windows/Fonts/consola.ttf', 24)

def shot(src, dst):
    lines = src.read_text(encoding='utf-8').splitlines()
    w, h = 3000, max(600, 135 + len(lines) * 34)
    im = Image.new('RGB',(w,h),'#161b22'); d=ImageDraw.Draw(im)
    d.rectangle((0,0,w,95),fill='#24292f'); d.text((40,28),src.name,font=font,fill='#ffffff')
    for i,line in enumerate(lines):
        d.text((35,115+i*34),f'{i+1:>2}',font=font,fill='#8b949e')
        d.text((105,115+i*34),line,font=font,fill='#e6edf3')
    dst.parent.mkdir(parents=True,exist_ok=True); im.save(dst)

doc=Document(); sec=doc.sections[0]; sec.orientation=WD_ORIENT.LANDSCAPE; sec.page_width,sec.page_height=Inches(11.69),Inches(8.27); sec.left_margin=sec.right_margin=Inches(.55); sec.top_margin=sec.bottom_margin=Inches(.5)
p=doc.add_paragraph(style='Title'); p.alignment=WD_ALIGN_PARAGRAPH.CENTER; p.add_run('AutoLoc Controllers et Tests Postman')
p=doc.add_paragraph('Ce document présente les neuf contrôleurs REST et la collection Postman CRUD associée. Chaque contrôleur délègue les opérations à la couche Service.'); p.alignment=WD_ALIGN_PARAGRAPH.CENTER
doc.add_page_break()
for n in names:
    doc.add_heading(n+'Controller',0)
    src=root/'src/main/java/tn/esprit/autoloc/controllers'/(n+'Controller.java'); dst=shots/(n+'Controller.png'); shot(src,dst)
    p=doc.add_paragraph(); p.alignment=WD_ALIGN_PARAGRAPH.CENTER; p.add_run().add_picture(str(dst),width=Inches(10.3))
    doc.add_paragraph('Routes Postman : GET liste, GET par id, POST, PUT et DELETE.')
    if n != names[-1]: doc.add_page_break()
doc.add_page_break(); doc.add_heading('Collection Postman CRUD',0)
doc.add_paragraph('La collection AutoLoc CRUD contient 45 requêtes : cinq opérations CRUD pour chacune des neuf entités. Configurez baseUrl avec http://localhost:8081 puis exécutez la collection dans Postman Runner.')
collection=root/'postman/AutoLoc CRUD.postman_collection.json'; dst=shots/'PostmanCollection.png'; shot(collection,dst)
p=doc.add_paragraph(); p.alignment=WD_ALIGN_PARAGRAPH.CENTER; p.add_run().add_picture(str(dst),width=Inches(10.3))
doc.save(out); print(out)
